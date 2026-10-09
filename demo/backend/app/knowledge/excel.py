"""Excel implementation of the knowledge source (first implementation; reread on every run)."""
from __future__ import annotations

import hashlib
import warnings
from datetime import datetime
from pathlib import Path
from typing import Any, Optional

from openpyxl import load_workbook

from .base import KnowledgeBase, KRow


class ExcelKnowledgeSource:
    """Reads every sheet whose header row starts with `id`.

    `columns` maps internal field names to the header text used in Maru's file, per sheet,
    e.g. {"Steel_prices": {"eur_per_kg": "Hind €/kg"}}. Missing mappings mean identical names.
    """

    def __init__(self, path: Path, columns: Optional[dict[str, dict[str, str]]] = None, name: str = "Maru knowledge file"):
        self.path = Path(path)
        self.columns = columns or {}
        self.name = name

    def load(self) -> KnowledgeBase:
        if not self.path.exists():
            raise FileNotFoundError(f"Knowledge file not found: {self.path}")
        raw = self.path.read_bytes()
        version = hashlib.sha256(raw).hexdigest()[:12]
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            wb = load_workbook(self.path, data_only=True, read_only=True)
        kb = KnowledgeBase(source_name=self.path.name, version=version, path=str(self.path))
        for ws in wb.worksheets:
            header: Optional[list[str]] = None
            header_row = 0
            rows: list[KRow] = []
            for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
                if header is None:
                    if row and str(row[0]).strip().lower() == "id":
                        header = [str(h).strip() if h is not None else "" for h in row]
                        header_row = r_idx
                    continue
                if not row or row[0] in (None, ""):
                    continue
                data = {header[i]: row[i] for i in range(min(len(header), len(row))) if header[i]}
                mapping = self.columns.get(ws.title, {})
                for internal, external in mapping.items():
                    if external in data:
                        data[internal] = data[external]
                kid = str(row[0]).strip()
                kr = KRow(id=kid, sheet=ws.title, row=r_idx, data=data)
                rows.append(kr)
                kb.by_id[kid] = kr
            if header is not None:
                kb.tables[ws.title] = rows
        wb.close()
        kb.loaded_at = datetime.now().isoformat(timespec="seconds")  # type: ignore[attr-defined]
        return kb


def knowledge_summary(kb: KnowledgeBase) -> dict[str, Any]:
    status_counts: dict[str, int] = {}
    for rows in kb.tables.values():
        for r in rows:
            st = str(r.get("status", "")) or "n/a"
            status_counts[st] = status_counts.get(st, 0) + 1
    return {
        "file": kb.source_name,
        "version": kb.version,
        "sheets": {k: len(v) for k, v in kb.tables.items()},
        "status_counts": status_counts,
        "price_lists": kb.price_lists(),
    }
