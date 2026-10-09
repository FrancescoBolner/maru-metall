"""Excel export with all data and references (one sheet per report section + every value)."""
from __future__ import annotations

from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..models import RunResult
from .view import TYPE_LABEL, build_view, ref_label, source_text

NAVY = "01437D"
TYPE_FILL = {"extracted": "E3EEF8", "calculated": "E6F2EC", "predicted": "FBEFD9"}
_FILLS = {k: PatternFill("solid", fgColor=c) for k, c in TYPE_FILL.items()}


def _sheet(wb: Workbook, title: str, header: list[str], rows: Iterable[list[Any]], widths: dict[int, int] | None = None,
           types: list[str | None] | None = None) -> None:
    ws = wb.create_sheet(title[:31])
    ws.append(header)
    hf, hfill, hal = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor=NAVY), Alignment(vertical="center", wrap_text=True)
    for c in ws[1]:
        c.font, c.fill, c.alignment = hf, hfill, hal
    for i, r in enumerate(rows):
        r = list(r)
        ws.append(r)
        if types and i < len(types) and types[i] in _FILLS:
            fill = _FILLS[types[i]]
            for col in range(1, len(r) + 1):
                ws.cell(row=i + 2, column=col).fill = fill
    ws.freeze_panes = "A2"
    for i, h in enumerate(header, 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(i, min(max(len(h) + 4, 12), 50))
    ws.auto_filter.ref = ws.dimensions


def write_excel(res: RunResult, out: Path) -> None:
    v = res.values
    view = build_view(res)
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    s = res.summary
    rows = [("Project", res.project_name), ("Client", res.company), ("Report", res.report_name), ("Revision", res.revision),
            ("Status", res.status), ("Project level", f"{res.filemap.level} — {res.filemap.level_reason}"),
            ("Quantities from", s.get("quantity_source")), ("Pricing method", s.get("method")),
            ("Total price €", v[s["price_id"]].value), ("Price range low €", v[s["low_id"]].value),
            ("Price range high €", v[s["high_id"]].value), ("Detailed method €", v[s["detailed_id"]].value),
            ("Quick €/t method €", v[s["quick_id"]].value), ("Steel kg", v["total.kg"].value),
            ("Working hours", v[s["hours_id"]].value), ("Workshop hours", v[s["metal_hours_id"]].value),
            ("Labour cost €", v[s["labour_cost_id"]].value), ("Paint litres", v[s["paint_litres_id"]].value),
            ("Painting cost €", v[s["paint_cost_id"]].value), ("CO2e t (placeholder factors)", v[s["co2_id"]].value),
            ("Share extracted", s["confidence"]["extracted"]), ("Share calculated", s["confidence"]["calculated"]),
            ("Share predicted", s["confidence"]["predicted"]), ("Time to report s", res.timings.get("total")),
            ("AI provider", res.ai.get("provider")), ("AI cost €", res.ai.get("cost_eur")),
            ("Knowledge file", f"{res.knowledge.get('file')} ({res.knowledge.get('version')})")]
    ws.append(["Item", "Value"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
    for r in rows:
        ws.append(list(r))
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 70

    _sheet(wb, "Offer", ["No.", "Name", "Surface", "Amount", "Unit", "Price €", "Unit price"],
           [[r["no"], r["name"], r["surface"], r["amount"], r["unit"], r["price"], r["unit_price"]] for r in view["offer_rows"]],
           {2: 60})
    _sheet(wb, "Categories", ["Category", "Weight kg", "Pieces", "Area m²", "Plate share", "Surface", "Surface type", "In scope", "Longest mm"],
           [[c.name, v[c.kg_id].value, v[c.pieces_id].value, v[c.area_id].value, v[c.plate_share_id].value, v[c.surface_id].value,
             v[c.surface_id].type, c.included, c.longest_mm] for c in res.categories], {1: 40})
    lt = [v[l.kg_id].type for l in res.lines]
    _sheet(wb, "Takeoff lines", ["Line", "Category", "Phase", "Profile", "Grade", "Plate", "Pieces", "Weight kg", "Length m", "Area m²",
                                 "Type", "Confidence", "Edited", "Reference", "Snippet"],
           [[l.id, l.category, l.phase, l.profile, l.grade, l.is_plate, v[l.pieces_id].value, v[l.kg_id].value,
             v[l.length_id].value if l.length_id else None, v[l.area_id].value, TYPE_LABEL[v[l.kg_id].type], v[l.kg_id].confidence,
             bool(v[l.kg_id].edited or v[l.pieces_id].edited), "; ".join(ref_label(r) for r in v[l.kg_id].refs[:3]),
             (v[l.kg_id].refs[0].snippet or "")[:200] if v[l.kg_id].refs else ""] for l in res.lines],
           {2: 26, 14: 50, 15: 60}, lt)
    feats = []
    ft = []
    for g, ids in res.features.items():
        for i in ids:
            x = v.get(i)
            if x:
                feats.append([g, x.label, str(x.value), x.unit, TYPE_LABEL[x.type], x.confidence, bool(x.conflict), source_text(x)])
                ft.append(x.type)
    _sheet(wb, "Features", ["Group", "Feature", "Value", "Unit", "Type", "Confidence", "Conflict", "Source"], feats, {2: 34, 3: 50, 8: 80}, ft)
    _sheet(wb, "Hours", ["Operation", "Category", "Quantity", "Unit", "Hours", "Norm", "Knowledge rows", "Costed in"],
           [[h.operation, h.category, v[h.quantity_id].value if h.quantity_id else None, v[h.quantity_id].unit if h.quantity_id else None,
             v[h.hours_id].value, h.norm, ", ".join(h.kn), h.costed_in] for h in res.hours], {2: 34, 6: 50})
    prows = []
    for p in res.paint:
        if not p.litres:
            prows.append([p.category, p.system, v[p.area_id].value, 0, "galvanised", None, v[p.hours_id].value, v[p.cost_id].value, ", ".join(p.kn)])
        for x in p.litres:
            prows.append([p.category, p.system, v[p.area_id].value, x.get("coats"), x["product"], v[x["value_id"]].value,
                          v[p.hours_id].value, v[p.cost_id].value, ", ".join(p.kn)])
    _sheet(wb, "Painting", ["Category", "System", "Area m²", "Coats", "Product", "Litres", "Painting hours", "Cost €", "Knowledge"], prows, {1: 34})
    _sheet(wb, "Material", ["Profile", "Grade", "Category", "Phase", "Net kg", "Waste factor", "€/kg", "Cost €", "Reusable offcuts kg", "Price row", "Waste rule"],
           [[m["profile"], m["grade"], m["category"], m["phase"], m["kg"], m["factor"], m["eur_per_kg"], v[m["cost_id"]].value,
             round(m.get("reusable_kg", 0), 1), m["price_kn"], v[m["waste_id"]].formula] for m in res.material.get("lines", [])],
           {11: 70})
    crow = [[c["group"], c["label"], v[c["value_id"]].value, ", ".join(c.get("kn") or []), v[c["value_id"]].formula]
            for c in res.cost.get("lines", [])]
    for x in res.cost.get("extras", []):
        crow.append(["Extras", x["label"], v[x["sale_id"]].value, "", v[x["sale_id"]].formula])
    crow.append(["Transport", res.cost.get("transport_info", {}).get("region"), v[res.cost["transport"]].value, "",
                 v[res.cost["transport"]].formula])
    crow.append(["Margin", "", v[res.cost["margin"]].value, "GEN-03, GEN-04", v[res.cost["margin"]].formula])
    crow.append(["TOTAL", s["method"], v[s["price_id"]].value, "", v[s["price_id"]].formula])
    _sheet(wb, "Cost", ["Group", "Line", "€", "Knowledge", "Formula"], crow, {2: 40, 5: 80})
    _sheet(wb, "Quick method", ["Category", "€/t", "Price €", "Hours", "Reference project", "Knowledge row", "Formula"],
           [[q["category"], v[q["eur_t_id"]].value, v[q["sale_id"]].value, v[q["hours_id"]].value, q["ref_project"], q["kn"],
             v[q["eur_t_id"]].formula] for q in res.cost.get("quick_lines", [])], {1: 34, 7: 90})
    car = res.carbon
    crows = [[g["label"], g["tonnes"], v[g["share_id"]].value, g.get("required"), g["recycled_planned"], g["compliant"], v[g["co2_id"]].value]
             for g in car.get("groups", [])]
    _sheet(wb, "Carbon", ["Steel group", "Bought t", "EAF share", "Required recycled", "Planned recycled", "Compliant", "t CO2e"], crows, {1: 30})
    ws2 = wb["Carbon"]
    ws2.append([])
    for i in (car.get("steel_id"), car.get("paint_id"), car.get("galv_id"), car.get("transport_id"), car.get("workshop_id"), car.get("total_id")):
        if i:
            ws2.append([v[i].label, None, None, None, None, None, v[i].value, v[i].formula])
    ws2.append([])
    for o in car.get("options", []):
        ws2.append([o["label"], None, None, None, None, None, v[o["delta_id"]].value])
    ws2.append([])
    for fl in car.get("flags", []):
        ws2.append([fl["flag"], fl["detail"], fl["reference"]])
    ws2.append(["Factors are placeholders until supplier EPDs (EN 15804) are available."])
    _sheet(wb, "Risks", ["#", "Severity", "Kind", "Title", "Detail", "Question", "References"],
           [[r.rank, r.severity, r.kind, r.title, r.detail, r.question, "; ".join(ref_label(x) for x in r.refs[:4])] for r in res.risks],
           {4: 40, 5: 80, 6: 50, 7: 50})
    _sheet(wb, "Questions", ["#", "Question", "Why"], [[q.id, q.text, q.reason] for q in res.questions], {2: 80, 3: 60})
    _sheet(wb, "Assumptions", ["Kind", "Text"], [[a["kind"], a.get("text")] for a in res.assumptions], {2: 120})
    _sheet(wb, "File map", ["Path", "Type", "Size bytes", "Pages/sheets/elements", "Date", "Status", "Reason", "Decided by", "Read by",
                            "Duplicate of", "Superseded by", "SHA-256", "Error"],
           [[e.path, e.kind, e.size, e.pages or e.sheets or e.elements, e.date, e.status, e.reason, e.decided_by, e.read_by,
             e.duplicate_of, e.superseded_by, e.sha256, e.error] for e in res.filemap.entries], {1: 70, 7: 70})
    vt = []
    vrows = []
    for x in v.values():
        vt.append(x.type)
        vrows.append([x.id, x.label, x.value if not isinstance(x.value, (list, dict)) else str(x.value), x.unit, TYPE_LABEL[x.type],
                      x.confidence, x.formula, ", ".join(x.kn), x.reasoning, x.question,
                      " | ".join(f"{ref_label(r)}: {(r.snippet or '')[:120]}" for r in x.refs[:4]),
                      f"{x.edited.before} → {x.edited.after} ({x.edited.reason or x.edited.source})" if x.edited else None,
                      " vs ".join(str(o.value) for o in x.conflict) if x.conflict else None])
    _sheet(wb, "All values", ["Id", "Label", "Value", "Unit", "Type", "Confidence", "Formula", "Knowledge rows", "Reasoning",
                              "Question", "References", "Edited", "Conflict"], vrows, {1: 30, 2: 44, 7: 60, 9: 60, 11: 90}, vt)
    if res.changes:
        _sheet(wb, "Changes", ["Change", "Before", "After", "Unit", "Reason / source"],
               [[c["label"], str(c["before"]), str(c["after"]), c.get("unit"), c.get("reason") or c.get("source")] for c in res.changes], {1: 60})
    wb.save(out)
